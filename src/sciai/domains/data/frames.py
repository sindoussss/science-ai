"""Reading datasets and applying recipes. Runs inside the tool sandbox.

A dataset descriptor (a node result of kind ``dataset``) names an imported file by its
SHA-256, the options it is read with, and a recipe of operations (filter, derive) applied
on top. Frames are rebuilt from that description on demand (no derived data is stored),
and kept in a small per-process cache keyed by the descriptor's key.

Reading rules, recorded in the options so a re-read is identical:
- the header is the first row; an empty header cell becomes ``column<N>``; duplicate
  names are refused (pandas would silently rename them);
- exactly the tokens in ``NA_TOKENS`` mean "missing" (pandas' longer default list is off);
- rows where every cell is missing are dropped;
- CSV encoding is UTF-8 (with or without BOM), else Windows-1252, else Latin-1.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

NA_TOKENS = ("", "NA", "N/A", "NaN", "nan", "null", "NULL", "None", "#N/A")
FORMATS = {".csv": "csv", ".txt": "csv", ".tsv": "tsv", ".tab": "tsv", ".xlsx": "xlsx", ".xlsm": "xlsx"}
MAX_LEVELS = 12          # text columns with at most this many distinct values list their levels
MAX_LEVEL_CHARS = 40
SNIFF_BYTES = 64 * 1024
CACHE_SIZE = 4
_UNIT = re.compile(r"^(?P<base>.*?\S)\s*[\[(](?P<unit>[^\[\]()]{1,24})[\])]\s*$")
# Bracketed words pint would read as units but a data column almost never means that way
# ("points" is a typographic length to pint, "a.u." arbitrary units, not astronomical ones).
NOT_UNITS = {"point", "points", "pt", "pts", "a.u.", "au", "arb. units", "count", "counts", "n", "score", "unit",
             "units", "items", "%ile", "rank"}
FILTER_OPS = ("==", "!=", "<", "<=", ">", ">=", "in", "not_in", "is_missing", "not_missing")


class DataError(ValueError):
    """A problem with the data or with how a tool asked for it; the message is shown to the model."""


# ------------------------------------------------------------------ hashing and keys
def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def source_key(sha256: str, fmt: str, options: dict[str, Any]) -> str:
    """The dataset id: the same bytes read the same way."""
    blob = json.dumps({"sha256": sha256, "format": fmt, "options": options}, sort_keys=True)
    return hashlib.sha256(blob.encode()).hexdigest()


def descriptor_key(source: dict[str, Any], recipe: list[dict[str, Any]]) -> str:
    blob = json.dumps({"source": source_key(source["sha256"], source["format"], source["options"]),
                       "recipe": recipe}, sort_keys=True)
    return hashlib.sha256(blob.encode()).hexdigest()


def format_of(path: str | Path) -> str:
    fmt = FORMATS.get(Path(path).suffix.lower())
    if fmt is None:
        raise DataError(f"{Path(path).name}: unsupported file type; use .csv, .tsv or .xlsx")
    return fmt


# ------------------------------------------------------------------ options
def _encoding(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ("utf-8", "cp1252"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def sniff_options(path: str | Path, fmt: str) -> dict[str, Any]:
    path = Path(path)
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = wb.sheetnames[0]
        finally:
            wb.close()
        return {"sheet": sheet, "na": list(NA_TOKENS)}
    encoding = _encoding(path)
    with open(path, encoding=encoding, newline="") as fh:
        sample = fh.read(SNIFF_BYTES)
    if fmt == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    decimal = "."
    if delimiter == ";" and len(re.findall(r"(?<![\d.,])-?\d+,\d+(?![\d.,])", sample)) > \
            len(re.findall(r"(?<![\d.,])-?\d+\.\d+(?![\d.,])", sample)):
        decimal = ","
    return {"delimiter": delimiter, "decimal": decimal, "encoding": encoding, "na": list(NA_TOKENS)}


# ------------------------------------------------------------------ reading
def clean_names(raw: list[Any]) -> list[str]:
    names = []
    for i, value in enumerate(raw):
        name = "" if value is None else " ".join(str(value).split())
        names.append(name or f"column{i + 1}")
    dupes = sorted({n for n in names if names.count(n) > 1})
    if dupes:
        raise DataError(f"duplicate column names: {', '.join(dupes)}; rename them and import again")
    return names


def _csv_header(path: Path, options: dict[str, Any]) -> list[str]:
    with open(path, encoding=options["encoding"], newline="") as fh:
        row = next(csv.reader(fh, delimiter=options["delimiter"]), None)
    if row is None:
        raise DataError("the file is empty")
    return clean_names(row)


def _xlsx_header(path: Path, sheet: str) -> list[str]:
    """The first row of the sheet, without the empty cells after the last named column."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        row = list(next(wb[sheet].iter_rows(values_only=True), ()))
    finally:
        wb.close()
    while row and (row[-1] is None or str(row[-1]).strip() == ""):
        row.pop()
    if not row:
        raise DataError("the sheet is empty")
    return clean_names(row)


def read_source(path: str | Path, fmt: str, options: dict[str, Any], max_rows: int) -> Any:
    """The file as a DataFrame, read exactly as ``options`` say."""
    import pandas as pd

    path = Path(path)
    na = list(options.get("na", NA_TOKENS))
    if fmt == "xlsx":
        names = _xlsx_header(path, options["sheet"])
        # dtype=object keeps the cell values as openpyxl gives them; _settle_types then types the
        # columns after empty rows are gone (pandas would turn True/False/empty into 1.0/0.0/NaN).
        df = pd.read_excel(path, sheet_name=options["sheet"], header=0, names=names, usecols=list(range(len(names))),
                           na_values=na, keep_default_na=False, nrows=max_rows + 1, engine="openpyxl", dtype=object)
    else:
        names = _csv_header(path, options)
        df = pd.read_csv(path, sep=options["delimiter"], decimal=options["decimal"], encoding=options["encoding"],
                         header=0, names=names, na_values=na, keep_default_na=False, nrows=max_rows + 1,
                         skip_blank_lines=True, engine="c", quotechar='"', index_col=False)
    df = _settle_types(df.dropna(how="all").reset_index(drop=True))
    if len(df) > max_rows:
        raise DataError(f"the file has more than {max_rows:,} rows; raise [data] max_rows to load it")
    return df


def is_missing(v: Any) -> bool:
    """A missing cell as it comes out of a frame row: None, NaN, pandas' NA or NaT."""
    if v is None:
        return True
    if isinstance(v, float):
        return math.isnan(v)
    return type(v).__name__ in ("NAType", "NaTType")


def _settle_types(df: Any) -> Any:
    """Type the columns pandas left as objects: all true/false values become a boolean column
    (with missing cells allowed), all numbers a numeric one, all dates a datetime one."""
    import datetime as dt

    import numpy as np
    import pandas as pd

    for name in df.columns:
        s = df[name]
        if s.dtype != object:
            continue
        vals = [v for v in s.tolist() if not is_missing(v)]
        if not vals:
            df[name] = s.astype(float)
        elif all(isinstance(v, (bool, np.bool_)) for v in vals):
            df[name] = s.astype("boolean")
        elif all(isinstance(v, str) and v.strip().lower() in ("true", "false") for v in vals):
            df[name] = s.map(lambda v: None if is_missing(v) else v.strip().lower() == "true").astype("boolean")
        elif all(isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, (bool, np.bool_))
                 for v in vals):
            df[name] = pd.to_numeric(s.where(s.notna(), None).astype(float) if len(vals) < len(s) else s)
        elif all(isinstance(v, (dt.datetime, dt.date, pd.Timestamp)) for v in vals):
            df[name] = pd.to_datetime(s)
    return df


def column_type(series: Any) -> str:
    import pandas as pd

    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_integer_dtype(series):
        return "integer"
    if pd.api.types.is_numeric_dtype(series):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "datetime"
    return "text"


def unit_of(name: str) -> str | None:
    m = _UNIT.match(name)
    if not m:
        return None
    unit = m.group("unit").strip()
    if unit.lower() in NOT_UNITS:
        return None
    try:
        from sciai.domains.physics import quantities as Q

        Q.make(1, unit)
    except Exception:  # noqa: BLE001 - not a unit pint knows: "(n)", "[mean]"...
        return None
    return unit


def schema_of(df: Any) -> list[dict[str, Any]]:
    out = []
    for name in df.columns:
        s = df[name]
        col: dict[str, Any] = {"name": str(name), "type": column_type(s), "unit": unit_of(str(name)),
                               "missing": int(s.isna().sum())}
        if col["type"] in ("text", "boolean"):
            present = s.dropna().astype(str)
            levels = present.unique()
            # Only categories are listed: few distinct short values, each used about twice or more.
            # A free-text column (notes, ids) shows no values, since the model sees no row data.
            if len(levels) <= MAX_LEVELS and 2 * len(levels) <= len(present) and \
                    all(len(v) <= MAX_LEVEL_CHARS for v in levels):
                col["levels"] = sorted(levels.tolist())
        out.append(col)
    return out


def checks_of(df: Any) -> dict[str, Any]:
    """Per-column figures an independent re-read must reproduce: numeric sums, true counts."""
    out: dict[str, Any] = {}
    for name in df.columns:
        s = df[name]
        t = column_type(s)
        if t == "boolean":
            out[str(name)] = {"true": int(s.fillna(False).astype(bool).sum())}
        elif t in ("integer", "number"):
            vals = [float(v) for v in s.dropna().tolist()]
            out[str(name)] = {"sum": math.fsum(vals), "count": len(vals)}
    return out


# ------------------------------------------------------------------ descriptors and cache
_cache: "OrderedDict[str, Any]" = OrderedDict()


def _cached(key: str, build: Any) -> Any:
    if key in _cache:
        _cache.move_to_end(key)
        return _cache[key]
    value = build()
    _cache[key] = value
    while len(_cache) > CACHE_SIZE:
        _cache.popitem(last=False)
    return value


def clear_cache() -> None:
    _cache.clear()


def require_descriptor(data: Any, what: str = "dataset") -> dict[str, Any]:
    if not isinstance(data, dict) or data.get("kind") != "dataset" or "source" not in data:
        raise DataError(f"{what} is not a dataset")
    return data


def load(desc: dict[str, Any], max_rows: int = 5_000_000, use_cache: bool = True) -> Any:
    """The descriptor's frame: the source file read, then its recipe applied. Returns a copy.
    ``use_cache=False`` (the UI's preview) reads without keeping the frame in memory."""
    desc = require_descriptor(desc)
    src = desc["source"]
    base_key = source_key(src["sha256"], src["format"], src["options"])

    def read() -> Any:
        if not Path(src["path"]).exists():
            raise DataError(f"the stored copy of {desc.get('name', 'the dataset')} is missing ({src['path']})")
        return read_source(src["path"], src["format"], src["options"], max_rows)

    recipe = list(desc.get("recipe") or [])
    if not use_cache:
        return apply_recipe(read(), recipe) if recipe else read()
    frame = _cached(base_key, read)
    if recipe:
        frame = _cached(descriptor_key(src, recipe), lambda: apply_recipe(frame, recipe))
    return frame.copy()


def make_descriptor(name: str, source: dict[str, Any], recipe: list[dict[str, Any]], df: Any,
                    with_checks: bool = False) -> dict[str, Any]:
    out = {"kind": "dataset", "name": name, "source": source, "recipe": recipe, "rows": int(len(df)),
           "columns": schema_of(df), "key": descriptor_key(source, recipe)}
    if with_checks:
        out["checks"] = checks_of(df)
    return out


def derived(desc: dict[str, Any], op: dict[str, Any], df: Any) -> dict[str, Any]:
    return make_descriptor(desc.get("name", "dataset"), desc["source"], [*desc.get("recipe", []), op], df)


# ------------------------------------------------------------------ columns
def resolve_column(desc_or_df: Any, name: str) -> str:
    """The real column name: exact, else a unique case-insensitive match, else a unique match of
    the name without its unit ("mass" for "mass [kg]"). Anything else is an error listing the columns."""
    if isinstance(desc_or_df, dict):
        cols = [c["name"] for c in desc_or_df.get("columns", [])]
        label = desc_or_df.get("name", "the dataset")
    else:
        cols = [str(c) for c in desc_or_df.columns]
        label = "the dataset"
    if name in cols:
        return name
    lower = [c for c in cols if c.lower() == str(name).lower()]
    if len(lower) == 1:
        return lower[0]
    base = [c for c in cols if (m := _UNIT.match(c)) and m.group("base").lower() == str(name).lower()]
    if len(base) == 1:
        return base[0]
    raise DataError(f"column {name!r} is not in {label}; its columns are: {', '.join(cols)}")


def numeric(df: Any, col: str) -> Any:
    import pandas as pd

    s = df[col]
    if pd.api.types.is_bool_dtype(s) or not pd.api.types.is_numeric_dtype(s):
        raise DataError(f"column {col!r} is {column_type(s)}, not numeric")
    return s.astype(float)


# ------------------------------------------------------------------ recipe operations
def _compare(series: Any, op: str, value: Any) -> Any:
    import numpy as np
    import pandas as pd

    if op == "is_missing":
        return series.isna()
    if op == "not_missing":
        return series.notna()
    numeric_col = pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series)
    if op in ("in", "not_in"):
        if not isinstance(value, list) or not value:
            raise DataError(f"{op} needs a non-empty list of values")
        vals = [float(v) for v in value] if numeric_col else [str(v) for v in value]
        target = series if numeric_col else series.astype(str)
        mask = target.isin(vals) & series.notna()
        return ~mask & series.notna() if op == "not_in" else mask
    if isinstance(value, (list, dict)) or value is None:
        raise DataError(f"{op} needs a single value")
    if numeric_col:
        try:
            v = float(value)
        except (TypeError, ValueError):
            raise DataError(f"column is numeric; {value!r} is not a number") from None
        target = series
    else:
        if op not in ("==", "!="):
            raise DataError(f"{op} needs a numeric column")
        v, target = str(value), series.astype(str)
    out = {"==": target == v, "!=": target != v, "<": target < v, "<=": target <= v,
           ">": target > v, ">=": target >= v}[op]
    return np.asarray(out, dtype=bool) & series.notna().to_numpy()


def apply_filter(df: Any, where: list[dict[str, Any]]) -> Any:
    import numpy as np

    mask = np.ones(len(df), dtype=bool)
    for cond in where:
        col = resolve_column(df, cond["column"])
        if cond["op"] not in FILTER_OPS:
            raise DataError(f"unknown filter op {cond['op']!r}; use one of {', '.join(FILTER_OPS)}")
        mask &= np.asarray(_compare(df[col], cond["op"], cond.get("value")), dtype=bool)
    return df[mask].reset_index(drop=True)


def derive_values(df: Any, expr: str, variables: dict[str, str]) -> Any:
    """``expr`` over columns, through the restricted parser; ``variables`` maps symbols to columns."""
    import numpy as np
    import sympy as sp

    from sciai.tools.parsing import parse, var

    e = parse(expr)
    unknown = sorted(str(s) for s in e.free_symbols if str(s) not in variables)
    if unknown:
        raise DataError(f"symbols {unknown} are not mapped to columns in vars")
    syms = [var(k) for k in variables]
    cols = [numeric(df, resolve_column(df, c)).to_numpy() for c in variables.values()]
    f = sp.lambdify(syms, e, modules="numpy")
    with np.errstate(all="ignore"):
        out = np.asarray(f(*cols), dtype=complex) * np.ones(len(df))
    real = np.where(np.abs(out.imag) < 1e-12, out.real, np.nan)
    return np.where(np.isfinite(real), real, np.nan)


def apply_recipe(df: Any, recipe: list[dict[str, Any]]) -> Any:
    for op in recipe:
        kind = op.get("op")
        if kind == "filter":
            df = apply_filter(df, op["where"])
        elif kind == "derive":
            if op["name"] in df.columns:
                raise DataError(f"column {op['name']!r} already exists")
            df = df.copy()
            df[op["name"]] = derive_values(df, op["expr"], op["vars"])
        else:
            raise DataError(f"unknown recipe operation {kind!r}")
    return df


# ------------------------------------------------------------------ independent re-reads
def _plain_float(text: str, decimal: str) -> float | None:
    t = text.strip()
    if decimal == ",":
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def reread_csv(path: str | Path, options: dict[str, Any]) -> dict[str, Any]:
    """Row/column counts, missing counts and numeric sums from Python's csv module, with no pandas."""
    na = set(options.get("na", NA_TOKENS))
    with open(path, encoding=options["encoding"], newline="") as fh:
        reader = csv.reader(fh, delimiter=options["delimiter"])
        header = clean_names(next(reader))
        cols: list[list[str]] = [[] for _ in header]
        rows = 0
        for row in reader:
            cells = [(row[i] if i < len(row) else "") for i in range(len(header))]
            if all(c.strip() in na or c in na for c in cells):
                continue
            rows += 1
            for i, c in enumerate(cells):
                cols[i].append(c)
    return _summarize(header, cols, rows, na, options.get("decimal", "."))


def reread_xlsx(path: str | Path, options: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    na = set(options.get("na", NA_TOKENS))
    header = _xlsx_header(Path(path), options["sheet"])
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb[options["sheet"]].iter_rows(values_only=True)
        next(it)
        cols: list[list[Any]] = [[] for _ in header]
        rows = 0
        for row in it:
            cells = [(row[i] if i < len(row) else None) for i in range(len(header))]
            if all(c is None or (isinstance(c, str) and c in na) for c in cells):
                continue
            rows += 1
            for i, c in enumerate(cells):
                cols[i].append(c)
    finally:
        wb.close()
    return _summarize(header, cols, rows, na, ".")


def _summarize(header: list[str], cols: list[list[Any]], rows: int, na: set[str], decimal: str) -> dict[str, Any]:
    missing, sums, trues = {}, {}, {}
    for name, values in zip(header, cols):
        present = [v for v in values if not (v is None or (isinstance(v, str) and (v in na or v.strip() in na)))]
        missing[name] = rows - len(present)
        nums = []
        for v in present:
            if isinstance(v, bool):
                nums = None
                break
            x = float(v) if isinstance(v, (int, float)) else _plain_float(str(v), decimal)
            if x is None:
                nums = None
                break
            nums.append(x)
        if nums is not None and present:
            sums[name] = {"sum": math.fsum(nums), "count": len(nums)}
        bools = [str(v).strip().lower() for v in present]
        if present and all(b in ("true", "false") for b in bools):
            trues[name] = sum(b == "true" for b in bools)
    return {"columns": header, "rows": rows, "missing": missing, "sums": sums, "trues": trues}


def _close(a: float, b: float) -> bool:
    return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))


def compare_reread(desc: dict[str, Any], reread: dict[str, Any]) -> list[str]:
    """Differences between what pandas loaded and what the independent re-read found."""
    problems = []
    names = [c["name"] for c in desc["columns"]]
    if names != reread["columns"]:
        problems.append(f"columns differ: {names} vs {reread['columns']}")
        return problems
    if desc["rows"] != reread["rows"]:
        problems.append(f"row count {desc['rows']} vs {reread['rows']}")
    for col in desc["columns"]:
        name = col["name"]
        if col["missing"] != reread["missing"].get(name):
            problems.append(f"{name}: {col['missing']} missing vs {reread['missing'].get(name)}")
    for name, figures in (desc.get("checks") or {}).items():
        if "sum" in figures:
            other = reread["sums"].get(name)
            if other is None or other["count"] != figures["count"] or not _close(other["sum"], figures["sum"]):
                problems.append(f"{name}: sum {figures['sum']!r} over {figures['count']} values vs {other}")
        if "true" in figures and reread["trues"].get(name) != figures["true"]:
            problems.append(f"{name}: {figures['true']} true vs {reread['trues'].get(name)}")
    return problems


def preview_rows(desc: dict[str, Any], n: int = 50) -> tuple[list[str], list[list[str]]]:
    """The first rows as text, for the UI's Data tab (read in the UI process, so not cached)."""
    df = load(desc, use_cache=False).head(n)
    rows = [["" if is_missing(v) else str(v) for v in r]
            for r in df.itertuples(index=False, name=None)]
    return [str(c) for c in df.columns], rows


def to_csv_text(df: Any) -> str:
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    return buf.getvalue()
